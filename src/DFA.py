from collections import defaultdict, deque
from dataclasses import dataclass


class State:
  def __init__(self, id: str, isStart: bool = False, isFinal: bool = False, transitions: list | None = None):
    self.id = id
    self.isStart = isStart
    self.isFinal = isFinal
    self.transitions = transitions if transitions is not None else []


class Transition:
  def __init__(self, from_state: State, to_state: State, symbol: str):
    self.from_state = from_state
    self.to_state = to_state
    self.symbol = symbol


class DFA:
  def __init__(self, start_state: State | None):
    self.start_state = start_state
    self.alphabet = []


  def from_transition_table(self, xlsx_path = "dfa-transitions.xlsx") -> None:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    ws = wb.active

    if ws is None:
      print("ws.active is None")
      return

    self.alphabet = [str(ws.cell(row=1, column=col).value) for col in range(3, ws.max_column + 1)]

    states: dict[str, State] = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
      state =  State(
        id=str(row[1].value),
        isStart=row[0].value != None and str(row[0].value).startswith("->"),
        isFinal=row[0].value != None and str(row[0].value).startswith("*")
      )
      states[state.id] = state
      if state.isStart:
        self.start_state = state

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
      state_id = str(row[1].value)
      state = states[state_id]
      transitions = []
      for col in range(2, ws.max_column):
        symbol = str(ws.cell(row=1, column=col + 1).value)
        to_state_id = str(row[col].value)
        if to_state_id != '-' and to_state_id != '' and to_state_id is not None:
          to_state = states[to_state_id]
          transition = Transition(from_state=state, to_state=to_state, symbol=symbol)
          transitions.append(transition)
      state.transitions = transitions

    sink = State("__SINK__")
    sink.transitions = [Transition(sink, sink, symbol) for symbol in self.alphabet]
    for state in states.values():
      existing = {t.symbol for t in state.transitions}
      for symbol in self.alphabet:
        if symbol not in existing:
          state.transitions.append(Transition(state, sink, symbol))


SYMBOL_ALIASES = {
  "space": " ",
  "tab": "\t",
  "new-line": "\n",
  "apostrophe": "'",
}


@dataclass(frozen=True)
class DFAMatch:
  end: int
  state: State


class DFAExecutor:
  def __init__(self, dfa: DFA):
    if dfa.start_state is None:
      raise ValueError("DFA must have a start state")
    self.dfa = dfa
    self.transitions: dict[State, dict[str, State]] = {}
    reverse = defaultdict(set)
    stack = [dfa.start_state]

    # Build transition lookup once
    while stack:
      state = stack.pop()
      if state in self.transitions:
        continue
      moves = {}
      for transition in state.transitions:
        symbol = SYMBOL_ALIASES.get(transition.symbol, transition.symbol)
        if symbol in moves and moves[symbol] is not transition.to_state:
          raise ValueError(f"Nondeterministic transition: {state.id}, {symbol!r}")
        moves[symbol] = transition.to_state
        reverse[transition.to_state].add(state)
        stack.append(transition.to_state)
      self.transitions[state] = moves

    # Find all states that can eventually reach a final state
    self.live_states = {state for state in self.transitions if state.isFinal}
    queue = deque(self.live_states)
    while queue:
      state = queue.popleft()
      for previous in reverse[state]:
        if previous not in self.live_states:
          self.live_states.add(previous)
          queue.append(previous)


  # def execute(self, input_string: str) -> bool:
  #   current_state = self.dfa.start_state
  #   for symbol in input_string:
  #     current_state = self.transitions[current_state].get(symbol)
  #     if current_state is None:
  #       return False
  #   return current_state.isFinal


  def match_longest(self, source: str, start: int = 0) -> DFAMatch | None:
    if not 0 <= start <= len(source):
      raise ValueError("Invalid starting position")

    current_state = self.dfa.start_state

    # Mohamed
    if current_state is None:
      return None;

    last_match = (DFAMatch(start, current_state) if current_state.isFinal else None)

    for i in range(start, len(source)):
      symbol = source[i]
      next_state = self.transitions[current_state].get(symbol)
      if next_state is None or next_state not in self.live_states:
        break
      current_state = next_state
      if current_state.isFinal:
        last_match = DFAMatch(i + 1, current_state)

    return last_match

